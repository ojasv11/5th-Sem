

import argparse
import csv
import re
from pathlib import Path

import openpyxl

DAYS = {"mon": "Monday", "tue": "Tuesday", "wed": "Wednesday",
        "thu": "Thursday", "fri": "Friday", "sat": "Saturday"}
TIME_RANGE = re.compile(r"^\s*(\d{1,2})[:.](\d{2})\s*[-\u2013]\s*(\d{1,2})[:.](\d{2})\s*$")

# Entries that never occupy a room: skipped, not reported as problems
SKIP = re.compile(r"^(lunch|elective|lib|break)$|online|swayam|training|extra-?curricular", re.I)

SUBJECT_CODE = re.compile(r"\b([A-Z]{2,5})\s*-\s*(\d{3})(?:\s*\(\s*[A-Z]\s*\))?")
GROUP_PREFIX = re.compile(r"^\s*(G\d)\s*[-\u2013]\s*", re.I)
LT_CR = re.compile(r"\b(LT|CR)\s*[-\u2013]?\s*(\d{1,2})\b", re.I)
LAB = re.compile(r"\b(?:([A-Za-z]+)\s+)?Lab\s*[-\u2013]?\s*(\d{1,2})\b", re.I)

OTHER_BUILDING = re.compile(r"\bblock\b|paramedical", re.I)
ROOM_WORDS = re.compile(r"\b(lab|floor|room|hall|block)\b", re.I)
SUBJECT_WITH_LAB = re.compile(r"\b[A-Z]{2,5}\s*-\s*\d{3}(?:\s*\(\s*[A-Z]\s*\))?(?:\s+Lab\b)?", re.I)

# Column names match the MySQL staging table, so the import wizard maps them automatically
ROW_FIELDS = ["day", "start_time", "end_time", "room", "subject",
              "course", "year", "section", "group_name", "source_cell"]
PROBLEM_FIELDS = ["Source Cell", "Day", "Time", "Text", "Issue"]


# ----------------------------------------------------------------- helpers
def clean(value):
    """Cell text with non-breaking spaces and line breaks collapsed."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def to_24h(hour, minute):
    """The sheet writes afternoon times as 1:00, 2:50 ... Classes run 8 am - 6 pm,
    so any hour below 8 is an afternoon hour."""
    hour = int(hour)
    if hour < 8:
        hour += 12
    return f"{hour:02d}:{int(minute):02d}"


def load_aliases(path):
    aliases = {}
    if path.exists():
        with open(path, newline="", encoding="utf-8-sig") as handle:
            for row in csv.DictReader(handle):
                aliases[clean(row["raw_name"]).lower()] = clean(row["room_name"])
    return aliases


def find_room(text, aliases):
    """Returns (room_name, text_without_room, issue).
    room None and issue None means: no room written, so this cell is a break."""
    lowered = text.lower()
    for raw in sorted(aliases, key=len, reverse=True):      # longest alias first
        if raw in lowered:
            start = lowered.index(raw)
            return aliases[raw], (text[:start] + text[start + len(raw):]).strip(), None

    if OTHER_BUILDING.search(text):
        return "OTHER", text, None

    match = LT_CR.search(text)
    if match:
        room = f"{match.group(1).upper()}{int(match.group(2))}"
        return room, (text[:match.start()] + text[match.end():]).strip(), None

    match = LAB.search(text)
    if match:
        qualifier = (match.group(1) or "").lower()
        if qualifier and qualifier != "lab" and not SUBJECT_CODE.search(match.group(0)):
            raw = text[match.start():match.end()]
            return None, text, f"Unrecognised room '{raw}': add it to room_aliases.csv"
        return f"Lab{int(match.group(2))}", (text[:match.start()] + text[match.end():]).strip(), None

    # No room pattern. If something room-like is still written, a person must map it;
    # otherwise it is a subject without a room, which counts as a break.
    leftover = SUBJECT_WITH_LAB.sub("", text)
    if ROOM_WORDS.search(leftover):
        return None, text, "Unrecognised room name: add it to room_aliases.csv"
    return None, text, None


def parse_entry(text, aliases):
    """One grid entry -> list of (group, subject, room, issue).
    'G1-PCS-511 Lab New Lab-1 & G2-PCS-503 Lab New Lab-2' gives two pieces."""
    results = []
    for piece in [p.strip() for p in text.split("&") if p.strip()]:
        group = "All"
        match = GROUP_PREFIX.match(piece)
        if match:
            group = match.group(1).upper()
            piece = piece[match.end():]

        room, rest, issue = find_room(piece, aliases)
        code = SUBJECT_CODE.search(rest)
        if code:
            subject = f"{code.group(1)}-{code.group(2)}"
        else:
            subject = re.sub(r"\s+", " ", rest).strip(" -,") or None
            if subject is None and room and not issue:
                issue = "No subject found in this cell"
        results.append((group, subject, room, issue))
    return results


def parse_title(lines):
    """Best-effort course / year / section from the text above a timetable grid."""
    text = " ".join(lines)
    info = {}
    sem = re.search(r"\b([IVX]+)(?:st|nd|rd|th)?\.?\s*SEMESTER", text, re.I)
    if sem:
        roman = {"I": 1, "V": 5, "X": 10}
        letters = sem.group(1).upper()
        value = sum(roman[c] for c in letters)
        for a, b in zip(letters, letters[1:]):
            if roman[a] < roman[b]:
                value -= 2 * roman[a]
        info["year"] = (value + 1) // 2
    if re.search(r"B\.?\s*TECH", text, re.I):
        info["course"] = "B.Tech CSE" if re.search(r"COMPUTER SCIENCE", text, re.I) else "B.Tech"
    elif re.search(r"\bBCA\b", text, re.I):
        info["course"] = "BCA"
    elif re.search(r"B\.?\s*COM", text, re.I):
        info["course"] = "B.Com"
    sec = re.search(r"(?:SECTION\s*[:\-]?\s*|[-\u2013]\s*)([A-Z]{1,3}\d{1,2})\b", text)
    if sec:
        info["section"] = sec.group(1)
    return info


# ------------------------------------------------------------- the parser
def merged_lookup(ws):
    """Map each top-left cell of a merged range to (min_col, max_col, max_row)."""
    lookup = {}
    for rng in ws.merged_cells.ranges:
        lookup[(rng.min_row, rng.min_col)] = (rng.min_col, rng.max_col, rng.max_row)
    return lookup


def find_header_rows(ws, merged):
    """Rows with at least 5 'H:MM-H:MM' cells. For each, returns
    times     {anchor_col: (start, end)}
    col_slot  {any_col_under_a_header_cell: anchor_col}  (headers can be merged, e.g. B:C)"""
    headers = []
    for row in ws.iter_rows():
        times, col_slot = {}, {}
        for cell in row:
            match = TIME_RANGE.match(clean(cell.value))
            if match:
                h1, m1, h2, m2 = match.groups()
                times[cell.column] = (to_24h(h1, m1), to_24h(h2, m2))
                min_col, max_col, _ = merged.get((cell.row, cell.column),
                                                 (cell.column, cell.column, cell.row))
                for c in range(min_col, max_col + 1):
                    col_slot[c] = cell.column
        if len(times) >= 5:
            headers.append((row[0].row, times, col_slot))
    return headers


def entry_time(min_col, max_col, times, col_slot):
    """Time covered by a cell spanning columns min_col..max_col. A cell owns a period
    when it covers ALL of that period's header columns; if it fully covers none, it
    takes the period of its own first column. Returns None outside the grid."""
    covered = set(range(min_col, max_col + 1))
    owned = [anchor for anchor in times
             if {c for c, a in col_slot.items() if a == anchor} <= covered]
    if not owned:
        owned = [col_slot[min_col]] if min_col in col_slot else []
    if not owned:
        return None
    return times[min(owned)][0], times[max(owned)][1]


def convert_sheet(ws, aliases, overrides, section_names, rows, problems, skipped_other):
    merged = merged_lookup(ws)
    headers = find_header_rows(ws, merged)
    prev_end = 0

    for index, (header_row, times, col_slot) in enumerate(headers):
        block_end = headers[index + 1][0] - 1 if index + 1 < len(headers) else ws.max_row
        block_no = len(section_names["used"]) + 1

        # --- course / year / section for this block
        title_lines = [clean(c.value) for r in ws.iter_rows(min_row=prev_end + 1, max_row=header_row - 1)
                       for c in r if clean(c.value)]
        info = parse_title(title_lines)
        info.update({k: v for k, v in overrides.items() if v is not None})
        if section_names["list"]:
            if block_no <= len(section_names["list"]):
                info["section"] = section_names["list"][block_no - 1]
            else:
                info.pop("section", None)
        section_names["used"].append(info.get("section", "TBD"))
        meta = {k: info.get(k, "TBD") for k in ("course", "year", "section")}
        home = re.search(r"Room No\.?\s*:?\s*-?\s*([A-Za-z]+\s*-?\s*\d+)", " ".join(title_lines))
        label = f"block {block_no} (header row {header_row}" + (f", home room {home.group(1)})" if home else ")")
        print(f"  {label} -> section {meta['section']}")
        if "TBD" in meta.values():
            missing = [k for k, v in meta.items() if v == "TBD"]
            problems.append({"Source Cell": f"{ws.title}!A{header_row}", "Day": "", "Time": "",
                             "Text": label,
                             "Issue": f"Missing {', '.join(missing)}: pass --course/--year/--sections"})

        # --- day rows in column A
        day_starts = []
        for r in range(header_row + 1, block_end + 1):
            day_label = clean(ws.cell(r, 1).value)[:3].lower()
            if day_label in DAYS:
                day_starts.append((r, DAYS[day_label]))

        for i, (day_row, day) in enumerate(day_starts):
            if (day_row, 1) in merged:
                day_end = merged[(day_row, 1)][2]
            else:
                day_end = day_starts[i + 1][0] - 1 if i + 1 < len(day_starts) else day_row

            # Cells starting in the same column on this day form one entry
            # (e.g. 'TCS-514' with 'LT-6' in the cell below it).
            entries = {}
            for r in range(day_row, day_end + 1):
                for col in range(2, ws.max_column + 1):
                    text = clean(ws.cell(r, col).value)
                    if not text:
                        continue
                    min_col, max_col, _ = merged.get((r, col), (col, col, r))
                    entry = entries.setdefault(min_col, {"texts": [], "cell": None, "max_col": max_col})
                    entry["texts"].append(text)
                    entry["max_col"] = max(entry["max_col"], max_col)
                    entry["cell"] = entry["cell"] or ws.cell(r, col).coordinate

            for min_col, entry in sorted(entries.items()):
                text = " ".join(entry["texts"])
                if SKIP.search(text):
                    continue
                slot = entry_time(min_col, entry["max_col"], times, col_slot)
                source = f"{ws.title}!{entry['cell']}"
                if slot is None:
                    problems.append({"Source Cell": source, "Day": day, "Time": "", "Text": text,
                                     "Issue": "Cell is outside the time columns"})
                    continue
                start, end = slot
                for group, subject, room, issue in parse_entry(text, aliases):
                    if issue:
                        problems.append({"Source Cell": source, "Day": day, "Time": f"{start}-{end}",
                                         "Text": text, "Issue": issue})
                    elif room == "OTHER":
                        skipped_other.append(f"{source} {day} {start}-{end}: {text}")
                    elif room:
                        rows.append({"day": day, "start_time": start, "end_time": end, "room": room,
                                     "subject": subject, "course": meta["course"], "year": meta["year"],
                                     "section": meta["section"], "group_name": group, "source_cell": source})
                    # no room and no issue: a break, ignored by design
        prev_end = block_end


def convert_files(files, course=None, year=None, sections=None, aliases_path=None):
    """Convert timetable workbooks without the command line (used by the GUI's Upload screen).
    Returns (rows, problems, skipped_other, blocks): blocks lists the section name
    given to each timetable block found, in order."""
    aliases = load_aliases(Path(aliases_path) if aliases_path
                           else Path(__file__).with_name("room_aliases.csv"))
    overrides = {"course": course, "year": year}
    section_names = {"list": list(sections or []), "used": []}
    rows, problems, skipped_other = [], [], []
    for path in files:
        workbook = openpyxl.load_workbook(path, data_only=True)
        for ws in workbook.worksheets:
            convert_sheet(ws, aliases, overrides, section_names, rows, problems, skipped_other)
    return rows, problems, skipped_other, section_names["used"]


def write_csv(path, fields, records):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def main():
    parser = argparse.ArgumentParser(description="Convert timetable grid workbooks to one CSV.")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--course")
    parser.add_argument("--year", type=int)
    parser.add_argument("--sections", help="comma-separated section names, one per block, e.g. A,B,C")
    parser.add_argument("--out-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    sections = [x.strip() for x in args.sections.split(",")] if args.sections else []
    for path in args.files:
        print(f"Reading {path}")
    rows, problems, skipped_other, blocks = convert_files(args.files, args.course, args.year, sections)
    if sections and len(sections) != len(blocks):
        print(f"WARNING: {len(sections)} section names given but {len(blocks)} timetable blocks found")

    args.out_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.out_dir / "timetable_rows.csv", ROW_FIELDS, rows)
    write_csv(args.out_dir / "timetable_problems.csv", PROBLEM_FIELDS, problems)
    print(f"{len(rows)} timetable rows written to {args.out_dir / 'timetable_rows.csv'}")
    print(f"{len(problems)} problems written to {args.out_dir / 'timetable_problems.csv'}")
    print(f"{len(skipped_other)} classes skipped because they are outside CSIT:")
    for line in skipped_other:
        print("   ", line)


if __name__ == "__main__":
    main()
